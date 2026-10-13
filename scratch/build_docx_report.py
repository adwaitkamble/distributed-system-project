"""Build a professional, publication-quality Microsoft Word (.docx) Academic Project Report
with embedded 300 DPI architecture and workflow diagrams, academic title page, certificates,
formatted tables, and complete technical specifications.
"""

import os
import re
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn, nsdecls

DIAGRAMS_DIR = os.path.join("submission_package", "diagrams")
REPORT_MD_PATH = os.path.join("docs", "FINAL_PROJECT_REPORT.md")
OUTPUT_DOCX_PATH = os.path.join("submission_package", "Final_Project_Report.docx")

def set_cell_bg(cell, hex_color):
  """Set background color of a table cell."""
  shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
  cell._tc.get_or_add_tcPr().append(shading_elm)

def set_cell_padding(cell, top=120, bottom=120, left=160, right=160):
  """Set internal padding for a table cell (in dxa: 20 dxa = 1 pt)."""
  tcPr = cell._tc.get_or_add_tcPr()
  tcMar = OxmlElement('w:tcMar')
  for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
    node = OxmlElement(f'w:{m}')
    node.set(qn('w:w'), str(val))
    node.set(qn('w:type'), 'dxa')
    tcMar.append(node)
  tcPr.append(tcMar)

def add_callout_box(doc, text, title="NOTE"):
  """Add a shaded callout block."""
  tbl = doc.add_table(rows=1, cols=1)
  tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
  cell = tbl.cell(0, 0)
  set_cell_bg(cell, "F1F5F9")
  set_cell_padding(cell, top=140, bottom=140, left=200, right=200)
  
  # Left border thick accent
  tcPr = cell._tc.get_or_add_tcPr()
  tcBorders = parse_xml(f'<w:tcBorders {nsdecls("w")}><w:left w:val="single" w:sz="24" w:space="0" w:color="1E3A8A"/><w:top w:val="none"/><w:right w:val="none"/><w:bottom w:val="none"/></w:tcBorders>')
  tcPr.append(tcBorders)
  
  p = cell.paragraphs[0]
  p.paragraph_format.space_before = Pt(2)
  p.paragraph_format.space_after = Pt(2)
  run_t = p.add_run(f"[{title}] ")
  run_t.bold = True
  run_t.font.name = "Calibri"
  run_t.font.size = Pt(10)
  run_t.font.color.rgb = RGBColor(30, 58, 138)
  
  run = p.add_run(text)
  run.font.name = "Calibri"
  run.font.size = Pt(10)
  run.font.italic = True
  run.font.color.rgb = RGBColor(51, 65, 85)
  
  doc.add_paragraph().paragraph_format.space_after = Pt(4)

def add_code_block(doc, code_text):
  """Add a shaded code block with monospace font."""
  tbl = doc.add_table(rows=1, cols=1)
  tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
  cell = tbl.cell(0, 0)
  set_cell_bg(cell, "F8FAFC")
  set_cell_padding(cell, top=100, bottom=100, left=160, right=160)
  
  tcPr = cell._tc.get_or_add_tcPr()
  tcBorders = parse_xml(f'<w:tcBorders {nsdecls("w")}><w:left w:val="single" w:sz="8" w:space="0" w:color="CBD5E1"/><w:top w:val="single" w:sz="8" w:space="0" w:color="CBD5E1"/><w:right w:val="single" w:sz="8" w:space="0" w:color="CBD5E1"/><w:bottom w:val="single" w:sz="8" w:space="0" w:color="CBD5E1"/></w:tcBorders>')
  tcPr.append(tcBorders)
  
  p = cell.paragraphs[0]
  p.paragraph_format.space_before = Pt(2)
  p.paragraph_format.space_after = Pt(2)
  p.paragraph_format.line_spacing = 1.05
  run = p.add_run(code_text.strip())
  run.font.name = "Consolas"
  run.font.size = Pt(9.0)
  run.font.color.rgb = RGBColor(15, 23, 42)
  
  doc.add_paragraph().paragraph_format.space_after = Pt(4)

def add_figure_image(doc, img_filename, caption_text):
  """Insert centered high-resolution figure with caption."""
  img_path = os.path.join(DIAGRAMS_DIR, img_filename)
  if os.path.exists(img_path):
    p_img = doc.add_paragraph()
    p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_img.paragraph_format.space_before = Pt(12)
    p_img.paragraph_format.space_after = Pt(4)
    run_img = p_img.add_run()
    run_img.add_picture(img_path, width=Inches(6.2))
    
    p_cap = doc.add_paragraph()
    p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_cap.paragraph_format.space_before = Pt(2)
    p_cap.paragraph_format.space_after = Pt(12)
    run_cap = p_cap.add_run(caption_text)
    run_cap.font.name = "Calibri"
    run_cap.font.size = Pt(9.5)
    run_cap.font.italic = True
    run_cap.font.bold = True
    run_cap.font.color.rgb = RGBColor(71, 85, 105)
  else:
    print(f"Warning: Image file not found: {img_path}")

def format_table(tbl, col_widths, headers, rows):
  """Render an academic table with dark header and borders."""
  tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
  
  # Header row
  hdr_cells = tbl.rows[0].cells
  for idx, title in enumerate(headers):
    hdr_cells[idx].text = title.strip()
    set_cell_bg(hdr_cells[idx], "1E3A8A")
    set_cell_padding(hdr_cells[idx], top=140, bottom=140, left=140, right=140)
    p = hdr_cells[idx].paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for r in p.runs:
      r.font.name = "Calibri"
      r.font.size = Pt(9.5)
      r.font.bold = True
      r.font.color.rgb = RGBColor(255, 255, 255)
      
  # Data rows
  for r_idx, row_data in enumerate(rows):
    row_cells = tbl.add_row().cells
    bg_color = "F8FAFC" if (r_idx % 2 == 1) else "FFFFFF"
    for c_idx, val in enumerate(row_data):
      if c_idx < len(row_cells):
        row_cells[c_idx].text = val.strip()
        set_cell_bg(row_cells[c_idx], bg_color)
        set_cell_padding(row_cells[c_idx], top=100, bottom=100, left=140, right=140)
        p = row_cells[c_idx].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT if c_idx > 0 else WD_ALIGN_PARAGRAPH.CENTER
        for r in p.runs:
          r.font.name = "Calibri"
          r.font.size = Pt(9.0)
          r.font.color.rgb = RGBColor(30, 41, 59)
          
  # Set column widths
  for row in tbl.rows:
    for idx, width in enumerate(col_widths):
      if idx < len(row.cells):
        row.cells[idx].width = width

def build_cover_page(doc):
  """Build formal institutional title page."""
  p = doc.add_paragraph()
  p.alignment = WD_ALIGN_PARAGRAPH.CENTER
  p.paragraph_format.space_before = Pt(12)
  p.paragraph_format.space_after = Pt(2)
  r = p.add_run("PCCoE / SPPU — DEPARTMENT OF COMPUTER ENGINEERING")
  r.font.name = "Calibri"
  r.font.size = Pt(11)
  r.font.bold = True
  r.font.color.rgb = RGBColor(71, 85, 105)

  p_inst = doc.add_paragraph()
  p_inst.alignment = WD_ALIGN_PARAGRAPH.CENTER
  p_inst.paragraph_format.space_before = Pt(4)
  p_inst.paragraph_format.space_after = Pt(24)
  r_inst = p_inst.add_run("PIMPRI CHINCHWAD COLLEGE OF ENGINEERING, PUNE\n(An Autonomous Institute Affiliated to Savitribai Phule Pune University)")
  r_inst.font.name = "Calibri"
  r_inst.font.size = Pt(13)
  r_inst.font.bold = True
  r_inst.font.color.rgb = RGBColor(30, 58, 138)

  p_title = doc.add_paragraph()
  p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
  p_title.paragraph_format.space_before = Pt(18)
  p_title.paragraph_format.space_after = Pt(12)
  r_title = p_title.add_run("DISTRIBUTED FAULT-TOLERANT VOTING SYSTEM WITH PLUGGABLE LEADER ELECTION")
  r_title.font.name = "Calibri"
  r_title.font.size = Pt(20)
  r_title.font.bold = True
  r_title.font.color.rgb = RGBColor(15, 23, 42)

  p_sub = doc.add_paragraph()
  p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
  p_sub.paragraph_format.space_before = Pt(6)
  p_sub.paragraph_format.space_after = Pt(28)
  r_sub = p_sub.add_run("A CAPSTONE PROJECT REPORT\nSubmitted in partial fulfillment of the requirements for the Degree of\nBACHELOR OF TECHNOLOGY IN COMPUTER ENGINEERING")
  r_sub.font.name = "Calibri"
  r_sub.font.size = Pt(11)
  r_sub.font.italic = True
  r_sub.font.color.rgb = RGBColor(100, 116, 139)

  # Group members box
  tbl = doc.add_table(rows=1, cols=2)
  tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
  
  c1 = tbl.cell(0, 0)
  c1.width = Inches(3.3)
  set_cell_padding(c1, top=140, bottom=140, left=140, right=140)
  p1 = c1.paragraphs[0]
  p1.paragraph_format.line_spacing = 1.25
  r1_h = p1.add_run("PROJECT GROUP MEMBERS:\n")
  r1_h.bold = True
  r1_h.font.color.rgb = RGBColor(30, 58, 138)
  r1_body = p1.add_run("• Adwait Kamble (123B1D002)\n• Advait Nathe (123B1D029)\n• Nikhil Puppalwar (123B1D040)")
  r1_body.font.size = Pt(10.5)

  c2 = tbl.cell(0, 1)
  c2.width = Inches(3.2)
  set_cell_padding(c2, top=140, bottom=140, left=140, right=140)
  p2 = c2.paragraphs[0]
  p2.paragraph_format.line_spacing = 1.25
  r2_h = p2.add_run("PROJECT GUIDE:\n")
  r2_h.bold = True
  r2_h.font.color.rgb = RGBColor(30, 58, 138)
  r2_body = p2.add_run("Mr. Ganesh Deshmukh\nAssistant Professor\nDepartment of Computer Engineering")
  r2_body.font.size = Pt(10.5)

  p_foot = doc.add_paragraph()
  p_foot.alignment = WD_ALIGN_PARAGRAPH.CENTER
  p_foot.paragraph_format.space_before = Pt(36)
  p_foot.paragraph_format.space_after = Pt(0)
  r_foot = p_foot.add_run("ACADEMIC YEAR: 2025–2026\nPUNE, MAHARASHTRA, INDIA")
  r_foot.font.name = "Calibri"
  r_foot.font.size = Pt(11)
  r_foot.font.bold = True
  r_foot.font.color.rgb = RGBColor(71, 85, 105)

  doc.add_page_break()

def build_certificates_page(doc):
  """Build Certificate of Approval and Declaration page."""
  h1 = doc.add_heading(level=1)
  r1 = h1.add_run("Certificate of Approval")
  r1.font.color.rgb = RGBColor(30, 58, 138)

  cert_text = (
      "This is to certify that the project report entitled 'Distributed Fault-Tolerant Voting System "
      "with Pluggable Leader Election' submitted by Adwait Kamble [123B1D002], Advait Nathe [123B1D029], "
      "and Nikhil Puppalwar [123B1D040] is a bonafide record of work carried out by them under my guidance and "
      "supervision in partial fulfillment of the requirements for the Degree of Bachelor of Technology in "
      "Computer Engineering of Savitribai Phule Pune University during the academic year 2025–2026."
  )
  p_cert = doc.add_paragraph(cert_text)
  p_cert.paragraph_format.line_spacing = 1.2
  p_cert.paragraph_format.space_after = Pt(24)

  # Signatures table
  tbl = doc.add_table(rows=2, cols=2)
  tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
  signatures = [
      ("Mr. Ganesh Deshmukh\nProject Guide\nDept. of Computer Engineering", "Dr. [HOD Placeholder]\nHead of Department\nDept. of Computer Engineering"),
      ("_________________________\nInternal Examiner", "_________________________\nExternal Examiner")
  ]
  for r_i, sig_pair in enumerate(signatures):
    row_cells = tbl.rows[r_i].cells
    for c_i, sig_text in enumerate(sig_pair):
      row_cells[c_i].width = Inches(3.2)
      set_cell_padding(row_cells[c_i], top=160, bottom=160, left=80, right=80)
      p = row_cells[c_i].paragraphs[0]
      p.paragraph_format.line_spacing = 1.15
      p.paragraph_format.space_before = Pt(8)
      r = p.add_run(sig_text)
      r.font.name = "Calibri"
      r.font.size = Pt(10)
      r.font.bold = True
      r.font.color.rgb = RGBColor(30, 41, 59)

  doc.add_paragraph().paragraph_format.space_after = Pt(16)

  h2 = doc.add_heading(level=2)
  r2 = h2.add_run("Candidate Declaration")
  r2.font.color.rgb = RGBColor(30, 58, 138)

  decl_text = (
      "We hereby declare that the project work presented in this report entitled 'Distributed Fault-Tolerant "
      "Voting System with Pluggable Leader Election' is an authentic record of our own research and engineering "
      "implementation carried out under the supervision of Mr. Ganesh Deshmukh. Any work, data, or concepts "
      "derived from existing academic literature or open-source software have been duly attributed and referenced."
  )
  p_decl = doc.add_paragraph(decl_text)
  p_decl.paragraph_format.line_spacing = 1.2
  p_decl.paragraph_format.space_after = Pt(16)

  p_names = doc.add_paragraph("Adwait Kamble [123B1D002]       Advait Nathe [123B1D029]       Nikhil Puppalwar [123B1D040]")
  p_names.paragraph_format.space_before = Pt(12)
  for r in p_names.runs:
    r.font.bold = True
    r.font.color.rgb = RGBColor(15, 23, 42)

  doc.add_page_break()

def parse_and_build_report():
  doc = docx.Document()

  # Set page margins to 1 inch
  for section in doc.sections:
    section.top_margin = Inches(1.0)
    section.bottom_margin = Inches(1.0)
    section.left_margin = Inches(1.0)
    section.right_margin = Inches(1.0)

  # Build Front Matter
  build_cover_page(doc)
  build_certificates_page(doc)

  # Read Markdown source
  with open(REPORT_MD_PATH, "r", encoding="utf-8") as f:
    lines = f.readlines()

  in_code_block = False
  code_buffer = []
  in_table = False
  table_headers = []
  table_rows = []

  # Skip front matter sections (1 & 2) since we rendered them professionally
  current_section_num = 0
  skip_until_section_3 = True

  for line in lines:
    stripped = line.strip()

    # Code block boundary
    if stripped.startswith("```"):
      if in_code_block:
        in_code_block = False
        add_code_block(doc, "\n".join(code_buffer))
        code_buffer = []
      else:
        in_code_block = True
        code_buffer = []
      continue

    if in_code_block:
      code_buffer.append(line.rstrip())
      continue

    # Table boundary
    if stripped.startswith("|") and stripped.endswith("|"):
      parts = [p.strip() for p in stripped.split("|")[1:-1]]
      # Check if delimiter row
      if all(set(p).issubset({"-", ":", " "}) for p in parts):
        continue
      if not in_table:
        in_table = True
        table_headers = parts
        table_rows = []
      else:
        table_rows.append(parts)
      continue
    else:
      if in_table:
        in_table = False
        # Render table
        col_count = len(table_headers)
        if col_count > 0:
          tbl = doc.add_table(rows=1, cols=col_count)
          widths = [Inches(6.2 / col_count)] * col_count
          # Specific custom widths for known tables
          if col_count == 6:  # Topology port mapping
            widths = [Inches(0.6), Inches(1.3), Inches(1.3), Inches(1.0), Inches(1.0), Inches(1.0)]
          elif col_count == 3:
            widths = [Inches(1.8), Inches(2.2), Inches(2.2)]
          format_table(tbl, widths, table_headers, table_rows)
          doc.add_paragraph().paragraph_format.space_after = Pt(4)
        table_headers = []
        table_rows = []

    # Section headers detection
    sec_match = re.match(r"^### Section (\d+):\s*(.*)", stripped)
    if sec_match:
      current_section_num = int(sec_match.group(1))
      sec_title = sec_match.group(2).strip()

      if current_section_num < 3:
        continue  # Already in custom front matter
      skip_until_section_3 = False

      # Page break before major sections
      if current_section_num in [3, 4, 5, 7, 9, 11, 14, 16, 20, 21, 23, 24, 25]:
        doc.add_page_break()

      h = doc.add_heading(level=1)
      h.paragraph_format.space_before = Pt(14)
      h.paragraph_format.space_after = Pt(6)
      r = h.add_run(f"Section {current_section_num}: {sec_title}")
      r.font.name = "Calibri"
      r.font.size = Pt(15)
      r.font.bold = True
      r.font.color.rgb = RGBColor(30, 58, 138)

      # Automatically insert figures at key sections
      if current_section_num == 9:
        add_figure_image(doc, "fig1_system_architecture.png", 
                         "Figure 1: Overall System Architecture — React SPA, Flask REST Bridges, NodeRuntime Engine, gRPC Communications, and Isolated MySQL Storage.")
      elif current_section_num == 11:
        add_figure_image(doc, "fig3_leader_election_failover.png", 
                         "Figure 3: Leader Election & Failover Sequence — Monotonic Heartbeat Expiration (3000ms), Bully/Ring Probes, Victory Broadcast, and Term Increment.")
      elif current_section_num == 14:
        add_figure_image(doc, "fig2_vote_replication_workflow.png", 
                         "Figure 2: Vote Submission & Two-Phase Replication Workflow — PENDING Staging, Concurrent gRPC Fan-out, Majority Quorum Verification, and Committed Persistence.")
      elif current_section_num == 16:
        add_figure_image(doc, "fig4_follower_recovery_sync.png", 
                         "Figure 4: Follower State Recovery & AutoFollowerSync Catch-up Protocol — Heartbeat Discovery, SyncLog Delta Catch-up, and 100% WAL Parity.")
      elif current_section_num == 20:
        add_figure_image(doc, "fig5_experimental_evaluation_harness.png", 
                         "Figure 5: Automated Benchmark Harness Architecture — Factorial Matrix Execution across EXP-01 to EXP-09, Continuous Telemetry, and Verifiable Artifact Generation.")
      continue

    if skip_until_section_3:
      continue

    # Subheadings
    if stripped.startswith("#### "):
      sub_title = stripped[5:].strip()
      h = doc.add_heading(level=2)
      h.paragraph_format.space_before = Pt(10)
      h.paragraph_format.space_after = Pt(4)
      r = h.add_run(sub_title)
      r.font.name = "Calibri"
      r.font.size = Pt(12)
      r.font.bold = True
      r.font.color.rgb = RGBColor(15, 23, 42)
      continue

    if stripped.startswith("##### "):
      sub_title = stripped[6:].strip()
      h = doc.add_heading(level=3)
      h.paragraph_format.space_before = Pt(8)
      h.paragraph_format.space_after = Pt(2)
      r = h.add_run(sub_title)
      r.font.name = "Calibri"
      r.font.size = Pt(11)
      r.font.bold = True
      r.font.color.rgb = RGBColor(51, 65, 85)
      continue

    # Blockquotes / Invariants
    if stripped.startswith("> "):
      quote_text = stripped[2:].strip()
      add_callout_box(doc, quote_text, title="INVARIANT / SPEC")
      continue

    # Horizontal rules
    if stripped in ["---", "***", "___"]:
      continue

    # Bullet lists
    if stripped.startswith("* ") or stripped.startswith("- "):
      bullet_text = stripped[2:].strip()
      p = doc.add_paragraph(style="List Bullet")
      p.paragraph_format.space_before = Pt(1)
      p.paragraph_format.space_after = Pt(2)
      p.paragraph_format.line_spacing = 1.15
      
      # Simple inline bold parsing: **text**
      tokens = re.split(r"(\*\*.*?\*\*)", bullet_text)
      for tok in tokens:
        if tok.startswith("**") and tok.endswith("**"):
          r = p.add_run(tok[2:-2])
          r.bold = True
          r.font.name = "Calibri"
          r.font.size = Pt(10.5)
        else:
          r = p.add_run(tok)
          r.font.name = "Calibri"
          r.font.size = Pt(10.5)
      continue

    # Numbered lists: 1. 2. 3.
    num_match = re.match(r"^(\d+)\.\s+(.*)", stripped)
    if num_match:
      num_str = num_match.group(1)
      item_text = num_match.group(2).strip()
      p = doc.add_paragraph(style="List Number")
      p.paragraph_format.space_before = Pt(1)
      p.paragraph_format.space_after = Pt(2)
      p.paragraph_format.line_spacing = 1.15
      
      tokens = re.split(r"(\*\*.*?\*\*)", item_text)
      for tok in tokens:
        if tok.startswith("**") and tok.endswith("**"):
          r = p.add_run(tok[2:-2])
          r.bold = True
          r.font.name = "Calibri"
          r.font.size = Pt(10.5)
        else:
          r = p.add_run(tok)
          r.font.name = "Calibri"
          r.font.size = Pt(10.5)
      continue

    # Regular paragraph
    if stripped:
      p = doc.add_paragraph()
      p.paragraph_format.space_before = Pt(2)
      p.paragraph_format.space_after = Pt(5)
      p.paragraph_format.line_spacing = 1.15
      
      tokens = re.split(r"(\*\*.*?\*\*)", stripped)
      for tok in tokens:
        if tok.startswith("**") and tok.endswith("**"):
          r = p.add_run(tok[2:-2])
          r.bold = True
          r.font.name = "Calibri"
          r.font.size = Pt(10.5)
        else:
          r = p.add_run(tok)
          r.font.name = "Calibri"
          r.font.size = Pt(10.5)

  # Save final document
  doc.save(OUTPUT_DOCX_PATH)
  print(f"Academic Report successfully generated at: {OUTPUT_DOCX_PATH}")

if __name__ == "__main__":
  parse_and_build_report()
